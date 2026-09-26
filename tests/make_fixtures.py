"""Generate the sample files in tests/fixtures/.

Run from the project folder:  python tests/make_fixtures.py

Every file is synthetic (no personal data), and contains known text so
tests can check that conversions kept the content. Re-running overwrites
the existing fixtures. The PDFs need LibreOffice installed.
"""
import io
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

FIXTURES = Path(__file__).parent / "fixtures"

# Known content that tests can assert on.
SIMPLE_HEADING = "Toolkit Sample Document"
SIMPLE_SENTENCE = "The quick brown fox jumps over the lazy dog."
SCANNED_TEXT = [
    "SCANNED PAGE",
    "This page is an image only.",
    "It has no extractable text layer.",
]
TABLE_ROWS = [
    ["Item", "Quantity", "Unit Price", "Total"],
    ["Notebook", "3", "2.50", "7.50"],
    ["Pen", "10", "0.80", "8.00"],
    ["Stapler", "1", "6.25", "6.25"],
    ["Folder", "5", "1.20", "6.00"],
]
PDF_PASSWORD = "test"


def find_soffice() -> str:
    found = shutil.which("soffice")
    if found:
        return found
    for env in ("ProgramFiles", "ProgramFiles(x86)"):
        base = os.environ.get(env)
        if base:
            candidate = Path(base) / "LibreOffice" / "program" / "soffice.exe"
            if candidate.exists():
                return str(candidate)
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if mac.exists():
        return str(mac)
    sys.exit("LibreOffice not found - it is needed to build the PDF fixtures.")


def to_pdf(src: Path, out_dir: Path) -> Path:
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            [find_soffice(), f"-env:UserInstallation={Path(profile).as_uri()}",
             "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)],
            check=True, capture_output=True, timeout=180,
        )
    return out_dir / (src.stem + ".pdf")


def sample_png() -> io.BytesIO:
    img = Image.new("RGB", (400, 200), "#dbe9f6")
    draw = ImageDraw.Draw(img)
    draw.rectangle([20, 20, 380, 180], outline="#1f5f99", width=4)
    draw.ellipse([150, 50, 250, 150], fill="#1f5f99")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def make_simple_docx(path: Path):
    doc = Document()
    doc.add_heading(SIMPLE_HEADING, level=1)
    doc.add_paragraph(SIMPLE_SENTENCE)
    doc.add_paragraph(
        "This second paragraph exists so conversions can be checked for "
        "paragraph breaks as well as text."
    )
    doc.save(path)


def make_complex_docx(path: Path):
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Local File Toolkit - Header Text"
    section.footer.paragraphs[0].text = "Confidential sample - Footer Text"

    doc.add_heading("Complex Sample Document", level=1)
    doc.add_paragraph("This document mixes several kinds of content.")

    doc.add_heading("Bullet list", level=2)
    for item in ["First bullet", "Second bullet", "Third bullet"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Numbered list", level=2)
    for item in ["Step one", "Step two", "Step three"]:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("Table", level=2)
    add_table(doc)

    doc.add_heading("Image", level=2)
    doc.add_picture(sample_png(), width=Inches(3))
    doc.add_paragraph("Figure 1: a generated sample image.")
    doc.save(path)


def add_table(doc):
    table = doc.add_table(rows=len(TABLE_ROWS), cols=len(TABLE_ROWS[0]))
    table.style = "Table Grid"  # visible borders help PDF table detection
    for r, row in enumerate(TABLE_ROWS):
        for c, value in enumerate(row):
            table.cell(r, c).text = value


def make_simple_xlsx(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    for row in TABLE_ROWS:
        ws.append(row)
    wb.save(path)


def make_multi_sheet_xlsx(path: Path):
    wb = Workbook()
    sales = wb.active
    sales.title = "Sales"
    sales.append(["Month", "Revenue"])
    for month, revenue in [("January", 1200), ("February", 1350), ("March", 990)]:
        sales.append([month, revenue])

    staff = wb.create_sheet("Staff")
    staff.append(["Name", "Team", "Start Year"])
    for row in [("Alex Doe", "Design", 2021), ("Sam Roe", "Support", 2023)]:
        staff.append(list(row))

    notes = wb.create_sheet("Notes")
    notes.append(["Note"])
    notes.append(["This sheet has a single text column."])
    wb.save(path)


def make_slides_pptx(path: Path):
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "Toolkit Sample Deck"
    title_slide.placeholders[1].text = "Three slides for conversion tests"

    for title, points in [
        ("Why local?", ["Files never leave your machine", "No uploads, no telemetry"]),
        ("Features", ["Convert to PDF", "Convert from PDF", "Merge PDFs", "Convert to Markdown"]),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = points[0]
        for point in points[1:]:
            body.add_paragraph().text = point
    prs.save(path)


def make_scanned_pdf(path: Path):
    """An image-only PDF, like a phone scan: no text layer at all."""
    page = Image.new("RGB", (1240, 1754), "white")  # A4 at 150 dpi
    draw = ImageDraw.Draw(page)
    try:
        font = ImageFont.truetype("arial.ttf", 48)
    except OSError:
        font = ImageFont.load_default(size=48)
    y = 200
    for line in SCANNED_TEXT:
        draw.text((150, y), line, fill="black", font=font)
        y += 90
    page = page.rotate(0.8, fillcolor="white")  # slight skew, like a real scan
    page.save(path, "PDF", resolution=150)


def make_protected_pdf(src: Path, path: Path):
    writer = PdfWriter(clone_from=str(src))
    writer.encrypt(PDF_PASSWORD)
    with open(path, "wb") as f:
        writer.write(f)


def main():
    FIXTURES.mkdir(parents=True, exist_ok=True)

    make_simple_docx(FIXTURES / "simple.docx")
    make_complex_docx(FIXTURES / "complex.docx")
    make_simple_xlsx(FIXTURES / "simple.xlsx")
    make_multi_sheet_xlsx(FIXTURES / "multi_sheet.xlsx")
    make_slides_pptx(FIXTURES / "slides.pptx")

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        shutil.copy(to_pdf(FIXTURES / "simple.docx", tmp), FIXTURES / "text.pdf")

        # tables.pdf comes from a Word table with borders, which PDF table
        # extractors detect far more reliably than an Excel sheet without gridlines.
        tables_docx = tmp / "tables.docx"
        doc = Document()
        doc.add_heading("Invoice Sample", level=1)
        add_table(doc)
        doc.save(tables_docx)
        shutil.copy(to_pdf(tables_docx, tmp), FIXTURES / "tables.pdf")

    make_scanned_pdf(FIXTURES / "scanned.pdf")
    make_protected_pdf(FIXTURES / "text.pdf", FIXTURES / "protected.pdf")
    (FIXTURES / "broken.pdf").write_bytes(b"this is not a pdf")
    shutil.copy(FIXTURES / "text.pdf", FIXTURES / "ümlaut ñame.pdf")

    for f in sorted(FIXTURES.iterdir()):
        print(f"  {f.name:<20} {f.stat().st_size:>8,} bytes")
    print(f"Done - fixtures written to {FIXTURES}")


if __name__ == "__main__":
    main()
